"""Produce every result of the project: CSV tables, results.json, README figures, and the planner workbook.

Usage: python -m demand.report   (after python -m demand.fetch)
"""
from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import LineChart, Reference  # noqa: E402
from openpyxl.formatting.rule import ColorScaleRule  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from scipy import stats  # noqa: E402

from demand import official, weekahead  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
TABLES = ROOT / "reports" / "tables"
FIG = ROOT / "reports" / "figures"
BLUE, GREY, RED = "#1f4e79", "#9aa5b1", "#b23a48"


def paired_test(a: pd.Series, b: pd.Series) -> dict:
    """Paired t-test on absolute percentage errors: is a's error different from b's?"""
    d = a - b
    t = stats.ttest_1samp(d, 0.0)
    return {"n": int(len(d)), "mean_difference_pp": round(float(d.mean()), 3), "p_value": round(float(t.pvalue), 4),
            "share_first_better_pct": round(float((a < b).mean() * 100), 1)}


def day_ahead_comparison(bt: pd.DataFrame) -> dict:
    g = bt[bt.h == 1].dropna(subset=["official"]).copy()
    g["combined"] = (g.ridge + g.official) / 2
    ape = {m: 100 * (g[m] - g.target).abs() / g.target for m in ("ridge", "official", "combined", "seasonal_naive")}
    return {
        "days": int(len(g)), "period": f"{g.target_date.min():%Y-%m-%d} to {g.target_date.max():%Y-%m-%d}",
        "mape_pct": {m: round(float(v.mean()), 2) for m, v in ape.items()},
        "ridge_vs_official": paired_test(ape["ridge"], ape["official"]),
        "combined_vs_official": paired_test(ape["combined"], ape["official"]),
    }


def figures(h: pd.DataFrame, audit: dict, scores: pd.DataFrame, combined_mape: float) -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    heat = audit["official_bias_hour_month"].set_index("hour")
    fig, ax = plt.subplots(figsize=(8, 5.2))
    lim = float(np.nanmax(np.abs(heat.values)))
    im = ax.imshow(heat.values, cmap="RdBu_r", vmin=-lim, vmax=lim, aspect="auto")
    ax.set_xticks(range(12), ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])
    ax.set_yticks(range(0, 24, 2), [f"{x:02d}:00" for x in range(0, 24, 2)])
    ax.set_title("Official day-ahead load forecast: mean error by hour and month, 2025–2026\n(red = forecast too high)")
    fig.colorbar(im, ax=ax, label="mean error, % of actual load")
    fig.tight_layout()
    fig.savefig(FIG / "official_bias_heatmap.png", dpi=150)
    plt.close(fig)

    s = audit["official_error_vs_solar"]
    fig, ax = plt.subplots(figsize=(7, 4))
    ax.bar(s.solar_bin.astype(str), s.mean_error_pct, color=[RED if v > 0 else BLUE for v in s.mean_error_pct])
    ax.axhline(0, color="black", lw=0.8)
    ax.set_xlabel("solar generation as % of load in that hour (10:00–16:00, 2023–2026)")
    ax.set_ylabel("mean forecast error, %")
    ax.set_title("The sunnier the hour, the more the official forecast over-predicts")
    fig.tight_layout()
    fig.savefig(FIG / "official_error_vs_solar.png", dpi=150)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7, 4))
    for model, colour, label in [("seasonal_naive", GREY, "same weekday last week"), ("ridge", BLUE, "calendar model (ridge)")]:
        m = scores[scores.model == model]
        ax.plot(m.horizon_days, m.mape_pct, marker="o", color=colour, label=label)
    off = scores[scores.model == "official"]
    ax.scatter(off.horizon_days - 0.08, off.mape_pct, color=RED, zorder=3, label=f"official day-ahead forecast ({off.mape_pct.iloc[0]:.2f} %)")
    ax.scatter([1 + 0.08], [combined_mape], color="#2e7d4f", marker="D", zorder=3, label=f"average of model and official ({combined_mape:.2f} %)")
    ax.set_xlabel("days ahead")
    ax.set_ylabel("mean absolute % error, daily demand")
    ax.set_ylim(0, None)
    ax.set_title("Week-ahead forecast accuracy, rolling backtest 2024–2026")
    ax.legend()
    fig.tight_layout()
    fig.savefig(FIG / "weekahead_accuracy.png", dpi=150)
    plt.close(fig)


def _header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def workbook(next_week: pd.DataFrame, bt: pd.DataFrame, scores: pd.DataFrame, audit: dict, path: Path) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Next 7 days"
    ws.append(["Forecast made", "Day", "Date", "Weekday", "Forecast (GWh)", "Low, 80 % (GWh)", "High, 80 % (GWh)", "Holiday", "Bridge day"])
    _header(ws)
    for r in next_week.itertuples():
        ws.append([r.origin.date(), r.h, r.target_date.date(), r.target_date.strftime("%A"), round(r.forecast_mwh / 1000, 1),
                   round(r.low_mwh / 1000, 1), round(r.high_mwh / 1000, 1), "yes" if r.holiday_national else "", "yes" if r.bridge_day else ""])
    n = len(next_week) + 1
    ws.append([])
    ws.append(["Total for the week", "", "", "", f"=SUM(E2:E{n})", "", ""])
    ws.cell(row=n + 2, column=1).font = Font(bold=True)
    ws.append(["Source: Bundesnetzagentur | SMARD.de (CC BY 4.0). Interval = 10th–90th percentile of backtest errors at that horizon."])
    chart = LineChart()
    chart.title = "Daily demand forecast with 80 % interval (GWh)"
    chart.y_axis.title = "GWh"
    chart.add_data(Reference(ws, min_col=5, max_col=7, min_row=1, max_row=n), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=4, min_row=2, max_row=n))
    chart.height, chart.width = 8, 18
    ws.add_chart(chart, "K2")
    _widths(ws, [14, 6, 12, 12, 14, 15, 15, 9, 10])
    for row in ws.iter_rows(min_row=2, max_row=n + 2, min_col=5, max_col=7):
        for c in row:
            c.number_format = "#,##0.0"

    ws = wb.create_sheet("Accuracy")
    ws.append(["Days ahead", "Model", "Days tested", "Mean abs. % error", "Bias %", "90th pct abs. % error"])
    _header(ws)
    names = {"seasonal_naive": "Same weekday last week", "ridge": "Calendar model (ridge)", "official": "Official day-ahead"}
    for r in scores.itertuples():
        ws.append([r.horizon_days, names[r.model], r.n, r.mape_pct, r.bias_pct, r.p90_abs_error_pct])
    _widths(ws, [11, 26, 12, 16, 9, 20])

    ws = wb.create_sheet("Backtest, day ahead")
    ws.append(["Date", "Actual (GWh)", "Calendar model (GWh)", "Official (GWh)", "Model abs. % error", "Official abs. % error", "Closer"])
    _header(ws)
    g = bt[bt.h == 1].dropna(subset=["official"]).sort_values("target_date")
    for i, r in enumerate(g.itertuples(), start=2):
        ws.append([r.target_date.date(), round(r.target / 1000, 2), round(r.ridge / 1000, 2), round(r.official / 1000, 2),
                   f"=ABS(C{i}-B{i})/B{i}", f"=ABS(D{i}-B{i})/B{i}", f'=IF(E{i}<F{i},"model","official")'])
        ws.cell(row=i, column=5).number_format = "0.00%"
        ws.cell(row=i, column=6).number_format = "0.00%"
    last = len(g) + 1
    ws.append([])
    ws.append(["Mean", "", "", "", f"=AVERAGE(E2:E{last})", f"=AVERAGE(F2:F{last})", f'=COUNTIF(G2:G{last},"model")/COUNTA(G2:G{last})'])
    for col in (5, 6, 7):
        ws.cell(row=last + 2, column=col).number_format = "0.00%"
        ws.cell(row=last + 2, column=col).font = Font(bold=True)
    ws.freeze_panes = "A2"
    _widths(ws, [12, 13, 20, 14, 18, 20, 10])

    ws = wb.create_sheet("Official bias by hour")
    heat = audit["official_bias_hour_month"]
    ws.append(["Hour"] + ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])
    _header(ws)
    for r in heat.itertuples(index=False):
        ws.append([f"{int(r[0]):02d}:00"] + [float(v) for v in r[1:]])
    ws.conditional_formatting.add("B2:M25", ColorScaleRule(start_type="min", start_color="2B6CB0", mid_type="num", mid_value=0,
                                                         mid_color="FFFFFF", end_type="max", end_color="C53030"))
    ws.append([])
    ws.append(["Mean % error of the official day-ahead forecast, 2025–2026. Positive = forecast above actual load."])
    _widths(ws, [8] + [7] * 12)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    h = official.load_hourly()
    audit = official.audit(h)
    corr = official.correction_results(h)
    for name, df in {**audit, **corr}.items():
        df.to_csv(TABLES / f"{name}.csv", index=False)
    daily = weekahead.load_daily()
    bt = weekahead.backtest(daily)
    scores = weekahead.scores(bt)
    nw = weekahead.next_week(daily, bt)
    bt.to_csv(TABLES / "weekahead_backtest.csv", index=False)
    scores.to_csv(TABLES / "weekahead_scores.csv", index=False)
    weekahead.holiday_scores(bt).to_csv(TABLES / "weekahead_holiday_scores.csv", index=False)
    nw.to_csv(TABLES / "weekahead_next_week.csv", index=False)
    cmp_ = day_ahead_comparison(bt)
    results = {
        "data_through": weekahead.LAST_FULL_DAY,
        "official_by_year": audit["official_by_year"].to_dict(orient="records"),
        "official_solar": audit["official_solar_correlation"].to_dict(orient="records")[0],
        "correction_daily_by_year": corr["correction_daily_by_year"].to_dict(orient="records"),
        "weekahead_scores": scores.to_dict(orient="records"),
        "holiday_scores": weekahead.holiday_scores(bt).to_dict(orient="records"),
        "day_ahead_comparison": cmp_,
    }
    (ROOT / "reports" / "results.json").write_text(json.dumps(results, indent=1, default=str), encoding="utf-8")
    figures(h, audit, scores, cmp_["mape_pct"]["combined"])
    workbook(nw, bt, scores, audit, ROOT / "reports" / "demand_forecast_planner.xlsx")
    print(json.dumps(cmp_, indent=1))


if __name__ == "__main__":
    main()
